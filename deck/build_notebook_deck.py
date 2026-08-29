"""Build a dark-theme workshop deck with rendered notebook screenshots."""

from __future__ import annotations

import sys
from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
NOTEBOOKS = ROOT / "notebooks"
ASSETS = Path(__file__).with_name("rendered_notebooks")
OUTPUT = Path(__file__).with_name("Context_Engineering_Healthcare_AI_Agents_Workshop.pptx")
sys.path.insert(0, str(NOTEBOOKS))

from workshop_utils import (  # noqa: E402
    ScenarioConfig,
    comparison_rows,
    expected_failure_evidence,
    phi_findings,
    run_agent,
)


BG = RGBColor(27, 27, 27)
PANEL = RGBColor(40, 40, 40)
CORAL = RGBColor(255, 95, 70)
WARM = RGBColor(250, 247, 242)
SOFT = RGBColor(205, 205, 205)
MUTED = RGBColor(145, 145, 145)
GREEN = RGBColor(80, 200, 145)


def font(size: int, bold: bool = False, mono: bool = False):
    candidates = [
        "/System/Library/Fonts/SFNSMono.ttf" if mono else "/System/Library/Fonts/SFNS.ttf",
        "/System/Library/Fonts/Monaco.ttf" if mono else "/System/Library/Fonts/Helvetica.ttc",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            try:
                return ImageFont.truetype(candidate, size=size)
            except OSError:
                pass
    return ImageFont.load_default()


def draw_wrapped(draw, text, xy, width_chars, fnt, fill, line_gap=8):
    x, y = xy
    for paragraph in text.splitlines() or [""]:
        lines = wrap(paragraph, width_chars, replace_whitespace=False) or [""]
        for line in lines:
            draw.text((x, y), line, font=fnt, fill=fill)
            y += fnt.size + line_gap
    return y


def notebook_image(filename: str, title: str, markdown: str, code: str, output: list[str]):
    ASSETS.mkdir(exist_ok=True)
    path = ASSETS / filename
    img = Image.new("RGB", (1600, 900), (246, 247, 249))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 1600, 72), fill=(31, 35, 41))
    d.text((32, 20), "databricks", font=font(26, True), fill=(255, 95, 70))
    d.text((205, 22), f"/Shared/context-engineering-healthcare-agents/notebooks/{title}", font=font(20), fill=(232, 234, 237))
    d.rectangle((0, 72, 1600, 122), fill=(255, 255, 255))
    d.text((34, 86), "Run all", font=font(18, True), fill=(36, 42, 48))
    d.text((132, 86), "Serverless", font=font(18), fill=(80, 87, 94))
    d.rounded_rectangle((28, 145, 1572, 300), radius=10, fill=(255, 255, 255), outline=(219, 222, 226), width=2)
    d.text((52, 164), title.replace("_", " "), font=font(27, True), fill=(25, 29, 33))
    draw_wrapped(d, markdown, (52, 210), 104, font(19), (70, 76, 82), 7)
    d.rounded_rectangle((28, 322, 1572, 600), radius=10, fill=(250, 251, 252), outline=(219, 222, 226), width=2)
    d.rectangle((28, 322, 38, 600), fill=(255, 95, 70))
    draw_wrapped(d, code, (58, 345), 118, font(18, mono=True), (36, 42, 48), 5)
    d.rounded_rectangle((28, 620, 1572, 864), radius=10, fill=(255, 255, 255), outline=(219, 222, 226), width=2)
    d.text((52, 642), "Output", font=font(18, True), fill=(80, 87, 94))
    y = 680
    for line in output:
        color = (202, 54, 54) if "BROKEN" in line or "PHI" in line else (30, 126, 85) if "FIXED" in line or "PASS" in line else (45, 50, 55)
        y = draw_wrapped(d, line, (52, y), 118, font(18, mono=True), color, 5) + 4
    img.save(path)
    return path


def add_text(slide, text, x, y, w, h, size=20, color=WARM, bold=False, font_name="DM Sans", align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    run = p.runs[0]
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def title(slide, text, kicker=None):
    if kicker:
        add_text(slide, kicker.upper(), 0.65, 0.34, 8.5, 0.28, 10, CORAL, True)
    add_text(slide, text, 0.65, 0.72, 12.0, 0.72, 27, WARM, True)


def footer(slide, n):
    add_text(slide, "DATABRICKS  |  CONTEXT ENGINEERING FOR HEALTHCARE AI AGENTS", 0.65, 7.17, 8.5, 0.18, 7, MUTED)
    add_text(slide, str(n), 12.15, 7.14, 0.4, 0.2, 8, MUTED, True, align=PP_ALIGN.RIGHT)


def add_bullets(slide, bullets, x=0.8, y=1.75, w=11.8, h=4.9, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(12)
        p.font.name = "DM Sans"
        p.font.size = Pt(size)
        p.font.color.rgb = SOFT
        p.text = "•  " + p.text


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    return slide


def screenshot_slide(prs, heading, caption, image_path, n):
    slide = blank(prs)
    title(slide, heading, "Notebook walkthrough")
    slide.shapes.add_picture(str(image_path), Inches(0.72), Inches(1.55), width=Inches(11.9), height=Inches(5.35))
    add_text(slide, caption, 0.82, 6.92, 10.9, 0.22, 8, MUTED)
    footer(slide, n)


def build():
    broken = run_agent(ScenarioConfig.broken())
    fixed = run_agent(ScenarioConfig.fixed())
    failures = expected_failure_evidence()
    comparison = comparison_rows(broken, fixed)

    setup_img = notebook_image(
        "00_setup.png", "00_Setup",
        "This workshop uses fictional patient markers only. Never paste real patient data into notebooks or MLflow traces.",
        "LIVE_MODEL = False\nBROKEN_CONFIG = ScenarioConfig.broken(live_model=LIVE_MODEL)\nFIXED_CONFIG = ScenarioConfig.fixed(live_model=LIVE_MODEL)",
        ["Inference mode: deterministic", "Safety fixture: synthetic identifiers only", "PASS — workshop configuration loaded"],
    )
    primer_img = notebook_image(
        "01_trace_primer.png", "01_Trace_Reading_Primer",
        "Open the trace and locate one span for each failure mode. Find the first bad span rather than starting with the final answer.",
        "primer_result = run_agent(ScenarioConfig.broken())\ndisplay(expected_failure_evidence())",
        [f"{f['failure'].upper():12}  {f['evidence']}" for f in failures],
    )
    broken_img = notebook_image(
        "02_broken_trace.png", "02_Combined_Lab — Broken Trace",
        "Generate a deliberately broken trace, follow synthetic PHI across boundaries, and diagnose retrieval and tool-selection failures.",
        "broken = run_agent(ScenarioConfig.broken(live_model=LIVE_MODEL))\ndisplay(phi_findings(broken))\nprint('Selected tool:', broken['selected_tool'])",
        [f"BROKEN — PHI locations: {len(phi_findings(broken))}", f"Selected tool: {broken['selected_tool']}", f"Retrieved chunks: {len(broken['retrieval'])}", "Memory retained raw sensitive answer"],
    )
    fixed_img = notebook_image(
        "03_before_after.png", "03_Before_After_Debrief",
        "Compare evidence rather than impressions. The fixed profile filters retrieval, guards context, narrows tools, and summarizes memory.",
        "fixed = run_agent(ScenarioConfig.fixed())\ndisplay(comparison_rows(broken, fixed))\nassert not phi_findings(fixed)",
        [f"FIXED — PHI locations: {len(phi_findings(fixed))}", f"Selected tool: {fixed['selected_tool']}", f"Retrieved chunks: {len(fixed['retrieval'])}", "PASS — safe memory and policy-only context"],
    )

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = blank(prs)
    add_text(slide, "CONTEXT ENGINEERING", 0.72, 0.65, 5.0, 0.3, 11, CORAL, True)
    add_text(slide, "Context Engineering for\nHealthcare AI Agents", 0.72, 1.35, 11.6, 1.65, 34, WARM, True)
    add_text(slide, "Diagnosing and tuning agent behavior from MLflow traces", 0.76, 3.38, 10.5, 0.45, 18, SOFT)
    add_text(slide, "Hands-on workshop  •  45 minutes  •  Fictional patient markers only", 0.76, 4.02, 10.5, 0.35, 13, MUTED)
    footer(slide, 1)

    slide = blank(prs); title(slide, "Context engineering optimizes the system—not one prompt", "Why it matters")
    add_bullets(slide, [
        "Curate the system prompt, retrieved evidence, tools, history, intermediate state, and memory available at inference.",
        "Treat context as finite: relevance, authority, recency, and minimization determine agent quality.",
        "In healthcare, the same boundary governs both answer quality and exposure of protected information.",
    ], size=22); footer(slide, 2)

    slide = blank(prs); title(slide, "Trace the complete context lifecycle", "Operating model")
    stages = [("SOURCE", 0.7), ("RETRIEVAL\n+ TOOLS", 3.25), ("CONTEXT\nGUARD", 5.95), ("MODEL", 8.7), ("MEMORY", 11.0)]
    for label, x in stages:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.35), Inches(1.65), Inches(1.05))
        shape.fill.solid(); shape.fill.fore_color.rgb = PANEL; shape.line.color.rgb = CORAL
        add_text(slide, label, x, 2.61, 1.65, 0.5, 15, WARM, True, align=PP_ALIGN.CENTER)
    for x in [2.48, 5.18, 7.93, 10.23]: add_text(slide, "→", x, 2.58, 0.55, 0.45, 25, CORAL, True, align=PP_ALIGN.CENTER)
    add_text(slide, "At each boundary: What entered? Was it authorized, relevant, current, and minimal? What was transformed, inferred, or persisted?", 1.1, 4.25, 11.1, 0.8, 19, SOFT, align=PP_ALIGN.CENTER)
    footer(slide, 3)

    screenshot_slide(prs, "Start with safe, deterministic workshop setup", "Rendered from 00_Setup.py; not a live workspace capture.", setup_img, 4)
    screenshot_slide(prs, "Read the first bad span—not merely the final answer", "Rendered from 01_Trace_Reading_Primer.py and deterministic workshop output.", primer_img, 5)

    slide = blank(prs); title(slide, "Four failure modes leave distinct trace evidence", "Trace diagnosis")
    cards = [("POISONING", "Untrusted context tries to alter agent behavior."), ("DISTRACTION", "Irrelevant volume crowds out useful signal."), ("CONFUSION", "Stale or contradictory evidence competes."), ("CLASH", "Overlapping tools or instructions collide.")]
    for i, (head, body) in enumerate(cards):
        x = 0.72 + (i % 2) * 6.2; y = 1.65 + (i // 2) * 2.25
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.72), Inches(1.75)); sh.fill.solid(); sh.fill.fore_color.rgb = PANEL; sh.line.color.rgb = PANEL
        add_text(slide, head, x+0.28, y+0.25, 5.1, 0.3, 14, CORAL, True); add_text(slide, body, x+0.28, y+0.75, 5.05, 0.65, 17, SOFT)
    footer(slide, 6)

    screenshot_slide(prs, "The broken trace exposes all four failure modes", "Rendered from 02_Combined_Lab.py and its deterministic broken scenario.", broken_img, 7)

    slide = blank(prs); title(slide, "Intervene at the earliest enforceable boundary", "Remediation profile")
    add_bullets(slide, [
        "Filter retrieval to approved policy corpora and current metadata; use smaller policy-oriented chunks.",
        "Drop unauthorized patient records and redact sensitive markers before model inference.",
        "Use progressive disclosure so policy intent exposes only the policy-search tool.",
        "Persist allowlisted, de-identified memory rather than raw conversations or model answers.",
    ], size=20); footer(slide, 8)

    screenshot_slide(prs, "Prove the fix with before-and-after evidence", "Rendered from 03_Before_After_Debrief.py and deterministic comparison output.", fixed_img, 9)

    slide = blank(prs); title(slide, "Evidence improves across every context boundary", "Before → after")
    headers = ["CHECK", "BROKEN", "FIXED"]
    rows = [[r["check"], str(r["broken"]), str(r["fixed"])] for r in comparison]
    table = slide.shapes.add_table(len(rows)+1, 3, Inches(0.85), Inches(1.6), Inches(11.65), Inches(4.7)).table
    table.columns[0].width = Inches(4.2); table.columns[1].width = Inches(3.5); table.columns[2].width = Inches(3.95)
    for c, value in enumerate(headers): table.cell(0,c).text = value
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row): table.cell(r,c).text = value
    for r in range(len(rows)+1):
        for c in range(3):
            cell = table.cell(r,c); cell.fill.solid(); cell.fill.fore_color.rgb = PANEL if r else CORAL
            for p in cell.text_frame.paragraphs:
                p.font.name = "DM Sans"; p.font.size = Pt(14); p.font.bold = r == 0; p.font.color.rgb = WARM
    footer(slide, 10)

    slide = blank(prs); title(slide, "45-minute participant journey", "Workshop flow")
    add_bullets(slide, [
        "0–3 — Setup and synthetic-data warning",
        "3–10 — Trace-reading primer and four failure modes",
        "10–25 — Generate and triage the broken trace",
        "25–39 — Apply remediation and compare trace evidence",
        "39–45 — Debrief, transfer rule, and exit check",
    ], size=21); footer(slide, 11)

    slide = blank(prs); title(slide, "Participant exit check", "Debrief")
    add_bullets(slide, [
        "Where did sensitive context first enter the trace?",
        "Which retrieval evidence proves poisoning, distraction, or confusion?",
        "Why did the wrong tool win, and how did progressive disclosure change routing?",
        "Which spans prove the fix changed the system rather than one answer?",
    ], size=21); footer(slide, 12)

    slide = blank(prs)
    add_text(slide, "THE DURABLE RULE", 0.72, 0.72, 4.0, 0.3, 11, CORAL, True)
    add_text(slide, "Trace the full lifecycle.\nCurate the minimum useful context.\nVerify the change with evidence.", 0.72, 1.55, 11.7, 2.6, 31, WARM, True)
    add_text(slide, "Workshop notebooks: /Shared/context-engineering-healthcare-agents", 0.76, 5.25, 10.5, 0.35, 13, MUTED)
    footer(slide, 13)

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
