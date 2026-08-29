"""Build a presentation focused on the 02_Combined_Lab notebook."""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "notebooks"))

from workshop_utils import ScenarioConfig, comparison_rows, phi_findings, run_agent  # noqa: E402


OUTPUT = Path(__file__).with_name("Combined_Lab_Healthcare_AI_Agents.pptx")
ASSETS = Path(__file__).with_name("rendered_notebooks")

BG = RGBColor(27, 27, 27)
PANEL = RGBColor(42, 42, 42)
CORAL = RGBColor(255, 95, 70)
WARM = RGBColor(250, 247, 242)
SOFT = RGBColor(208, 208, 208)
MUTED = RGBColor(148, 148, 148)
GREEN = RGBColor(80, 200, 145)


def add_text(slide, text, x, y, w, h, size=20, color=WARM, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    if align is not None:
        paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = "DM Sans"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    return slide


def heading(slide, title, kicker):
    add_text(slide, kicker.upper(), 0.65, 0.32, 9.0, 0.25, 10, CORAL, True)
    add_text(slide, title, 0.65, 0.70, 12.0, 0.72, 27, WARM, True)


def footer(slide, number):
    add_text(slide, "DATABRICKS  |  COMBINED LAB", 0.65, 7.16, 5.5, 0.18, 7, MUTED)
    add_text(slide, str(number), 12.15, 7.13, 0.4, 0.2, 8, MUTED, True, PP_ALIGN.RIGHT)


def bullets(slide, items, x=0.82, y=1.70, w=11.7, h=4.9, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {item}"
        paragraph.space_after = Pt(12)
        paragraph.font.name = "DM Sans"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = SOFT


def card(slide, x, y, w, h, label, body, accent=CORAL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = PANEL
    add_text(slide, label, x + 0.25, y + 0.20, w - 0.5, 0.28, 13, accent, True)
    add_text(slide, body, x + 0.25, y + 0.65, w - 0.5, h - 0.78, 16, SOFT)


def screenshot(slide, image, caption):
    slide.shapes.add_picture(str(image), Inches(0.75), Inches(1.52), width=Inches(11.82), height=Inches(5.32))
    add_text(slide, caption, 0.83, 6.89, 10.8, 0.20, 8, MUTED)


def build():
    broken = run_agent(ScenarioConfig.broken())
    fixed = run_agent(ScenarioConfig.fixed())
    comparison = comparison_rows(broken, fixed)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = blank(prs)
    add_text(slide, "HANDS-ON MLflow TRACE LAB", 0.72, 0.62, 6.0, 0.3, 11, CORAL, True)
    add_text(slide, "Combined Lab:\nDiagnose and Tune a Healthcare AI Agent", 0.72, 1.35, 11.8, 1.75, 33, WARM, True)
    add_text(slide, "Find synthetic PHI, context confusion, and tool clash—then prove the fix", 0.76, 3.50, 11.2, 0.48, 17, SOFT)
    add_text(slide, "27 minutes  •  Deterministic by default  •  Fictional workshop data only", 0.76, 4.20, 10.5, 0.35, 13, MUTED)
    footer(slide, 1)

    slide = blank(prs); heading(slide, "Your mission", "Combined lab")
    bullets(slide, [
        "Determine where synthetic PHI first enters the agent and where it crosses protected boundaries.",
        "Explain why retrieval underperforms and why the ambiguous router selects the patient-record tool.",
        "Apply five controlled context remediations without training or replacing the model.",
        "Compare two real nested MLflow traces and prove the system changed—not merely the wording of one answer.",
    ], size=20); footer(slide, 2)

    slide = blank(prs); heading(slide, "One task, three competing signals", "Broken scenario")
    add_text(slide, "“What follow-up care is needed after discharge for this patient?”", 0.9, 1.52, 11.5, 0.55, 20, WARM, True, PP_ALIGN.CENTER)
    card(slide, 0.72, 2.35, 3.85, 2.20, "CP-104 · ACTIVE", "Cardiology Discharge Policy\nFollow-up within 7 days")
    card(slide, 4.74, 2.35, 3.85, 2.20, "CM-220 · ACTIVE", "Care Management Standard\nFollow-up within 30 days")
    card(slide, 8.76, 2.35, 3.85, 2.20, "PATIENT TOOL", "Ambiguous router output\nFollow-up within 30 days")
    add_text(slide, "No retrieved precedence rule identifies the controlling policy.", 1.25, 5.20, 10.8, 0.45, 19, CORAL, True, PP_ALIGN.CENTER)
    footer(slide, 3)

    slide = blank(prs); heading(slide, "This is context confusion—not a knowledge gap", "Diagnosis")
    bullets(slide, [
        "The model receives mutually inconsistent authoritative evidence, plus a tool output reinforcing 30 days.",
        "The broken prompt correctly prevents outside clinical knowledge from silently breaking the tie.",
        "A plausible answer cannot repair missing policy hierarchy or metadata upstream.",
        "The safe behavior is to expose the unresolved conflict and fix the context before inference.",
    ], size=20); footer(slide, 4)

    slide = blank(prs); heading(slide, "Follow the first bad span", "Trace triage")
    stages = [
        ("RETRIEVE", "mixed policies\n+ patient note"),
        ("GUARD", "unsafe context\npasses through"),
        ("ROUTER", "overlapping tools\npatient tool wins"),
        ("MODEL", "PHI + conflict\nreach inference"),
        ("MEMORY", "raw transcript\nis persisted"),
    ]
    for i, (label, body) in enumerate(stages):
        x = 0.55 + i * 2.55
        card(slide, x, 2.20, 2.15, 2.15, label, body)
        if i < len(stages) - 1:
            add_text(slide, "→", x + 2.17, 2.95, 0.35, 0.4, 24, CORAL, True, PP_ALIGN.CENTER)
    add_text(slide, "Intervene at the earliest safe boundary—not only at final-answer rendering.", 0.9, 5.15, 11.5, 0.45, 19, WARM, True, PP_ALIGN.CENTER)
    footer(slide, 5)

    slide = blank(prs); heading(slide, "The broken run makes the defects observable", "Notebook output")
    screenshot(slide, ASSETS / "02_broken_trace.png", f"Deterministic run: {len(phi_findings(broken))} synthetic-PHI locations; selected tool: {broken['selected_tool']}.")
    footer(slide, 6)

    slide = blank(prs); heading(slide, "Five settings change the context lifecycle", "Remediation profile")
    items = [
        ("JUST-IN-TIME", "Task-specific retrieval timing is explicit in the trace."),
        ("SMALL CHUNKS", "Only approved policy chunks; exclude patient note and giant mixed chunk."),
        ("DROP + REDACT", "Unsafe documents are removed and synthetic PHI is redacted before inference."),
        ("PROGRESSIVE TOOLS", "Expose only search_clinical_policy for this policy-only intent."),
        ("SAFE MEMORY", "Persist an allowlisted, de-identified summary—not the raw transcript."),
    ]
    for i, (label, body) in enumerate(items):
        x = 0.72 + (i % 3) * 4.18
        y = 1.55 + (i // 3) * 2.25
        card(slide, x, y, 3.85, 1.85, label, body, GREEN)
    footer(slide, 7)

    slide = blank(prs); heading(slide, "Same agent workflow, safer configuration", "What is real")
    bullets(slide, [
        "run_agent executes root agent, retriever, guard, router, tool, model, and memory functions instrumented with mlflow.trace.",
        "Embedded documents, tools, and synthetic values are controlled fixtures—not production Vector Search, EHR, or MCP integrations.",
        "LIVE_MODEL=True calls databricks-claude-sonnet-4-5 with deterministic fallback; False simulates only the generated answer.",
        "ScenarioConfig.fixed() activates safer branches; it does not train or repair a model.",
    ], size=18); footer(slide, 8)

    slide = blank(prs); heading(slide, "The fixed run removes ambiguity before inference", "Notebook output")
    screenshot(slide, ASSETS / "03_before_after.png", f"Fixed run: {len(phi_findings(fixed))} synthetic-PHI locations; selected tool: {fixed['selected_tool']}.")
    footer(slide, 9)

    slide = blank(prs); heading(slide, "Compare evidence, not impressions", "Before → after")
    table = slide.shapes.add_table(len(comparison) + 1, 3, Inches(0.85), Inches(1.55), Inches(11.65), Inches(4.85)).table
    table.columns[0].width = Inches(4.20)
    table.columns[1].width = Inches(3.55)
    table.columns[2].width = Inches(3.90)
    for col, text in enumerate(["CHECK", "BROKEN", "FIXED"]):
        table.cell(0, col).text = text
    for row_index, row in enumerate(comparison, 1):
        for col, text in enumerate([row["check"], str(row["broken"]), str(row["fixed"])]):
            table.cell(row_index, col).text = text
    for row in range(len(comparison) + 1):
        for col in range(3):
            cell = table.cell(row, col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CORAL if row == 0 else PANEL
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "DM Sans"
                paragraph.font.size = Pt(14)
                paragraph.font.bold = row == 0
                paragraph.font.color.rgb = WARM
    footer(slide, 10)

    slide = blank(prs); heading(slide, "Exit check: prove each boundary changed", "Debrief")
    bullets(slide, [
        "retrieve_context no longer returns the patient record or giant mixed chunk.",
        "context_guard reports drop_and_redact with low PHI risk.",
        "tool_router exposes only the policy-search tool for this intent.",
        "model receives concise, de-identified, non-conflicting context.",
        "memory_write stores an allowlisted summary rather than the raw transcript.",
    ], size=19); footer(slide, 11)

    slide = blank(prs)
    add_text(slide, "THE DURABLE RULE", 0.72, 0.72, 4.0, 0.3, 11, CORAL, True)
    add_text(slide, "Find the first bad context boundary.\nFix it before inference.\nProve it in the trace.", 0.72, 1.55, 11.7, 2.55, 32, WARM, True)
    add_text(slide, "/Shared/context-engineering-healthcare-agents/notebooks/02_Combined_Lab", 0.76, 5.30, 11.5, 0.35, 13, MUTED)
    footer(slide, 12)

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
